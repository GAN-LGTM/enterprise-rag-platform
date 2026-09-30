"""
文档解析 —— 按扩展名路由到不同解析器（对应需求书 FR-KB-02）。

私有化环境常常装不全依赖，所以这里全部做成"可选依赖"：
  装了就精确解析（含页码/Sheet 名），没装则给出明确的安装提示，而不是静默失败。

生产推荐：unstructured.io（统一引擎）+ PaddleOCR（扫描件，本地 GPU/CPU）。
"""
from __future__ import annotations

from dataclasses import dataclass


@dataclass
class ParsedBlock:
    text: str
    page_num: int | None = None
    sheet: str | None = None


def parse_file(path: str, filename: str) -> list[ParsedBlock]:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "txt" or ext == "md":
        return _parse_text(path)
    if ext == "pdf":
        return _parse_pdf(path)
    if ext in ("docx", "doc"):
        return _parse_docx(path)
    if ext in ("xlsx", "xls", "csv"):
        return _parse_table(path, ext)
    raise ValueError(f"暂不支持的文件类型: .{ext}（支持 pdf/docx/xlsx/txt/md）")


def _parse_text(path: str) -> list[ParsedBlock]:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return [ParsedBlock(text=f.read())]


def _parse_pdf(path: str) -> list[ParsedBlock]:
    try:
        import pypdf  # noqa: PLC0415
    except ImportError as e:
        raise RuntimeError("解析 PDF 需要安装: pip install pypdf（生产建议 unstructured[pdf] + PaddleOCR）") from e

    out: list[ParsedBlock] = []
    with open(path, "rb") as f:
        reader = pypdf.PdfReader(f)
        for i, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                out.append(ParsedBlock(text=text, page_num=i))
    if not out:
        raise RuntimeError("PDF 无可提取文本（可能是扫描件），请开启 OCR：设置 OCR_ENABLED=true 并部署 PaddleOCR")
    return out


def _parse_docx(path: str) -> list[ParsedBlock]:
    try:
        import docx  # noqa: PLC0415
    except ImportError as e:
        raise RuntimeError("解析 Word 需要安装: pip install python-docx") from e

    d = docx.Document(path)
    blocks = [ParsedBlock(text=p.text.strip()) for p in d.paragraphs if p.text.strip()]
    # 表格：按行拼成 "列1 | 列2"，并保留表头语义
    for t in d.tables:
        rows = []
        header = None
        for r in t.rows:
            cells = [c.text.strip() for c in r.cells]
            if header is None:
                header = cells
                continue
            rows.append(" | ".join(f"{h}:{c}" for h, c in zip(header, cells)))
        if rows:
            blocks.append(ParsedBlock(text="\n".join(rows)))
    return blocks or [ParsedBlock(text="")]


def _parse_table(path: str, ext: str) -> list[ParsedBlock]:
    if ext == "csv":
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return [ParsedBlock(text=f.read())]
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError as e:
        raise RuntimeError("解析 Excel 需要安装: pip install openpyxl") from e

    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        lines = []
        for r in rows[1:]:
            if all(c is None for c in r):
                continue
            lines.append(" | ".join(f"{h}:{c}" for h, c in zip(header, r) if c is not None))
        if lines:
            out.append(ParsedBlock(text="\n".join(lines), sheet=ws.title))
    return out


# ------------------------------------------------------------
# OCR（扫描件）
# ------------------------------------------------------------
def ocr_pdf(path: str) -> list[ParsedBlock]:
    """PaddleOCR 本地识别。私有化部署时开启，不调用任何外部 API。"""
    try:
        from paddleocr import PaddleOCR  # noqa: PLC0415
        import fitz  # PyMuPDF，把 PDF 页转成图片
    except ImportError as e:
        raise RuntimeError("扫描件 OCR 需要: pip install paddleocr pymupdf") from e

    engine = PaddleOCR(use_angle_cls=True, lang="ch")
    doc = fitz.open(path)
    out = []
    for i, page in enumerate(doc, start=1):
        pix = page.get_pixmap(dpi=200)
        img_path = f"{path}.p{i}.png"
        pix.save(img_path)
        res = engine.ocr(img_path, cls=True)
        text = "\n".join(line[1][0] for line in (res[0] or []))
        if text:
            out.append(ParsedBlock(text=text, page_num=i))
    return out
